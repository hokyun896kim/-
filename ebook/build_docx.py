#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word(.docx) 전자책 생성기  ―  「나는 AI에게 종목을 묻지 않았다」 (원고 v1.1)
content.parse() 가 만든 요소 목록을 Word 네이티브 표/스타일로 렌더링한다.
산출물: ebook/build/book.docx
"""
import math, pathlib, re
from docx import Document
from docx.shared import Pt, Mm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import content as C

ROOTDIR = pathlib.Path(__file__).resolve().parent
BUILD = ROOTDIR / "build"; BUILD.mkdir(exist_ok=True)
IMG = BUILD / "img"
OUT = BUILD / "book.docx"

NAVY="163B4E"; NAVY2="21566E"; ACCENT="C0894B"; ACCENTD="9C6B33"
INK="232020"; MUTED="6E6A63"; LINE="DED8CC"
TINT="EEF3F2"; TINT2="F6F1E7"; GOODBG="EAF1ED"; BADBG="F7EEE8"; PROMPTBG="EEF3F2"
HEAD="맑은 고딕"; BODY="바탕"

# ---------- XML 헬퍼 ----------
def _set(el, tag, **a):
    e=OxmlElement(tag)
    for k,v in a.items(): e.set(qn(k), v)
    el.append(e); return e
def shade(cell,c):
    _set(cell._tc.get_or_add_tcPr(),'w:shd',**{'w:val':'clear','w:color':'auto','w:fill':c})
def cmar(cell,t=80,b=80,l=120,r=120):
    m=_set(cell._tc.get_or_add_tcPr(),'w:tcMar')
    for s,v in (('top',t),('bottom',b),('start',l),('end',r)):
        _set(m,f'w:{s}',**{'w:w':str(v),'w:type':'dxa'})
def borders(cell,color=LINE,sz=4,sides=('top','bottom','start','end')):
    bd=_set(cell._tc.get_or_add_tcPr(),'w:tcBorders')
    for s in sides: _set(bd,f'w:{s}',**{'w:val':'single','w:sz':str(sz),'w:space':'0','w:color':color})
def nobord(tbl):
    bd=_set(tbl._tbl.tblPr,'w:tblBorders')
    for s in ('top','left','bottom','right','insideH','insideV'):
        _set(bd,f'w:{s}',**{'w:val':'none','w:sz':'0','w:space':'0','w:color':'auto'})
def laccent(cell,color=NAVY,sz=26):
    bd=_set(cell._tc.get_or_add_tcPr(),'w:tcBorders')
    _set(bd,'w:start',**{'w:val':'single','w:sz':str(sz),'w:space':'0','w:color':color})
def vcenter(cell): _set(cell._tc.get_or_add_tcPr(),'w:vAlign',**{'w:val':'center'})

def run(p,text,font=BODY,size=10.5,bold=False,italic=False,color=INK):
    r=p.add_run(text); r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
    r.font.color.rgb=RGBColor.from_string(color); r.font.name=font
    rf=r._element.get_or_add_rPr().get_or_add_rFonts()
    rf.set(qn('w:ascii'),font); rf.set(qn('w:hAnsi'),font); rf.set(qn('w:eastAsia'),font)
    return r
def para(doc,align=None,before=0,after=5,line=1.5,indent=None):
    p=doc.add_paragraph(); pf=p.paragraph_format
    if align is not None: p.alignment=align
    pf.space_before=Pt(before); pf.space_after=Pt(after)
    pf.line_spacing=line; pf.line_spacing_rule=WD_LINE_SPACING.MULTIPLE
    if indent is not None: pf.left_indent=Mm(indent)
    return p
def botborder(p,color=ACCENT,sz=16):
    _set(_set(p._p.get_or_add_pPr(),'w:pBdr'),'w:bottom',
         **{'w:val':'single','w:sz':str(sz),'w:space':'6','w:color':color})
def box(doc,fill,accent=None):
    t=doc.add_table(rows=1,cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); shade(c,fill); cmar(c,110,110,150,150); borders(c,fill,2)
    if accent: laccent(c,accent,26)
    return c
def spacer(doc,pt=2): doc.add_paragraph().paragraph_format.space_after=Pt(pt)

# ---------- 컴포넌트 ----------
def keysentence(doc,text,icon=''):
    c=box(doc,TINT2,ACCENT)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.2
    run(p,"이 장의 핵심",font=HEAD,size=8.5,bold=True,color=ACCENTD)
    p2=c.add_paragraph(); p2.paragraph_format.space_after=Pt(0); p2.paragraph_format.line_spacing=1.45
    run(p2,text,font=HEAD,size=11,bold=True,color=NAVY); spacer(doc)

def opener(doc, icon, eyebrow, title, subtitle=None):
    # 그림이 있는 챕터 오프닝: 큰 일러스트 + (영문 eyebrow) + 제목 + 짧은 골드 선 + (부제)
    ip=IMG/f"hero_{icon}.png"
    p0=doc.add_paragraph(); p0.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.page_break_before=True; p0.paragraph_format.space_before=Pt(8); p0.paragraph_format.space_after=Pt(5)
    if icon and ip.exists(): p0.add_run().add_picture(str(ip), width=Mm(24))
    if eyebrow:
        pe=para(doc,WD_ALIGN_PARAGRAPH.CENTER,after=1); run(pe,eyebrow,font=HEAD,size=9,bold=True,color=ACCENT)
    h=doc.add_paragraph(style='Heading 2'); h.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run(h,title,font=HEAD,size=15,bold=True,color=NAVY)
    rr=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=2,after=(2 if subtitle else 7))
    rr.paragraph_format.left_indent=Mm(54); rr.paragraph_format.right_indent=Mm(54); botborder(rr,ACCENT,18)
    if subtitle:
        ps=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=1,after=7); run(ps,subtitle,font=HEAD,size=13,bold=True,color=NAVY)

def first_sentence_split(text):
    m=re.match(r'^([^"“”]{6,38}?[.?!])\s*(.*)$', text, re.S)
    if m and len(m.group(1))<=38: return m.group(1), m.group(2)
    return None, text

def ornament(doc):
    op=IMG/"ornament.png"
    p=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=3,after=4)
    if op.exists(): p.add_run().add_picture(str(op), width=Mm(34))
    else: run(p,"✦",color=ACCENT,size=11)

def body_para(doc,text,lead=False):
    p=para(doc,WD_ALIGN_PARAGRAPH.JUSTIFY)
    if lead and text:   # 장 첫 문단: 짧고 깔끔한 첫 문장만 굵게
        first,rest=first_sentence_split(text)
        if first:
            run(p,first,size=10.5,bold=True,color=NAVY); run(p," "+rest,size=10.5,color=INK)
        else:
            run(p,text,size=10.5,color=INK)
    else:
        run(p,text,size=10.5,color=INK)
def callout(doc,title,body):
    c=box(doc,TINT,NAVY)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.3
    run(p,"● ",size=8,bold=True,color=ACCENT,font=HEAD); run(p,title,font=HEAD,size=10.5,bold=True,color=NAVY)
    p2=c.add_paragraph(); p2.paragraph_format.space_after=Pt(0); p2.paragraph_format.line_spacing=1.4
    run(p2,body,size=9.8,color="3A4A46"); spacer(doc)
def compare(doc,headers,rows,good_right,caption=''):
    # 2열(라벨 열 제거) + 결론 캡션 — 모바일 가독성
    t=doc.add_table(rows=1+len(rows),cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for i,c in enumerate(t.rows[0].cells):
        shade(c,NAVY); cmar(c); vcenter(c)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
        run(p,headers[i+1],font=HEAD,size=9.8,bold=True,color="FFFFFF")
    bgA,bgB=(BADBG,GOODBG) if good_right else (GOODBG,BADBG)
    for ri,(lab,a,b) in enumerate(rows,1):
        cs=t.rows[ri].cells
        for c in cs: cmar(c)
        shade(cs[0],bgA); shade(cs[1],bgB)
        run(cs[0].paragraphs[0],a,size=9.6,color="3A3631"); run(cs[1].paragraphs[0],b,size=9.6,color="3A3631")
        for c in cs: c.paragraphs[0].paragraph_format.space_after=Pt(0); c.paragraphs[0].paragraph_format.line_spacing=1.3
    for row in t.rows:
        for c in row.cells: c.width=Mm(58)
    if caption:
        cp=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=2,after=0); run(cp,"→ "+caption,font=HEAD,size=9.4,bold=True,color=ACCENTD)
    spacer(doc)
def cards(doc,items):
    for i,(title,desc) in enumerate(items,1):
        t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; nobord(t)
        bc=t.cell(0,0); shade(bc,"ECE0C8"); vcenter(bc); cmar(bc,60,60,80,80); bc.width=Mm(8)
        bp=bc.paragraphs[0]; bp.alignment=WD_ALIGN_PARAGRAPH.CENTER; bp.paragraph_format.space_after=Pt(0)
        run(bp,str(i),font=HEAD,size=9,bold=True,color=ACCENTD)
        cc=t.cell(0,1); cmar(cc,40,40,150,40); cc.width=Mm(107)
        p=cc.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.2
        run(p,title,font=HEAD,size=10.5,bold=True,color=NAVY)
        p2=cc.add_paragraph(); p2.paragraph_format.space_after=Pt(0); p2.paragraph_format.line_spacing=1.3
        run(p2,desc,size=9.4,color="3A3631")
        spacer(doc,1)
def modebar(doc,modes):
    t=doc.add_table(rows=2,cols=len(modes)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; nobord(t)
    for i,(name,desc) in enumerate(modes):
        top=t.cell(0,i); shade(top,NAVY); cmar(top,60,40,30,30); vcenter(top)
        p=top.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(0)
        run(p,name,font=HEAD,size=9,bold=True,color="FFFFFF")
        bot=t.cell(1,i); shade(bot,"F3F0E9"); cmar(bot,40,60,20,20); borders(bot,LINE,2)
        p2=bot.paragraphs[0]; p2.alignment=WD_ALIGN_PARAGRAPH.CENTER; p2.paragraph_format.space_after=Pt(0); p2.paragraph_format.line_spacing=1.1
        run(p2,desc,font=HEAD,size=7.6,color="4A453F")
    w=Mm(116//len(modes))
    for row in t.rows:
        for c in row.cells: c.width=w
    spacer(doc)
def flow(doc,steps):
    for i,(title,desc) in enumerate(steps,1):
        t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; nobord(t)
        bc=t.cell(0,0); shade(bc,NAVY); vcenter(bc); cmar(bc,70,70,70,70); bc.width=Mm(10)
        bp=bc.paragraphs[0]; bp.alignment=WD_ALIGN_PARAGRAPH.CENTER; bp.paragraph_format.space_after=Pt(0)
        run(bp,str(i),font=HEAD,size=11,bold=True,color="FFFFFF")
        cc=t.cell(0,1); cmar(cc,40,40,150,40); cc.width=Mm(106); vcenter(cc)
        p=cc.paragraphs[0]; p.paragraph_format.space_after=Pt(1); p.paragraph_format.line_spacing=1.2
        run(p,title,font=HEAD,size=10.5,bold=True,color=NAVY)
        p2=cc.add_paragraph(); p2.paragraph_format.space_after=Pt(0); p2.paragraph_format.line_spacing=1.25
        run(p2,desc,size=9.3,color="3A3631")
        spacer(doc,1)
def dodont(doc,label,items,kind):
    fill=BADBG if kind=="dont" else GOODBG
    col="A6433B" if kind=="dont" else "2F6F5E"; mark="✕" if kind=="dont" else "✓"
    c=box(doc,fill,col)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(3); run(p,label,font=HEAD,size=10.5,bold=True,color=col)
    for t,d in items:
        pi=c.add_paragraph(); pi.paragraph_format.space_after=Pt(2); pi.paragraph_format.line_spacing=1.3
        run(pi,f"{mark} ",font=HEAD,size=9.5,bold=True,color=col)
        run(pi,t+" ",size=9.6,bold=True,color=NAVY); run(pi,"— "+d,size=9.4,color="4A453F")
    spacer(doc)

# ---------- 헤딩 ----------
def heading2(doc,text):
    h=doc.add_paragraph(style='Heading 2'); h.paragraph_format.page_break_before=True
    run(h,text,font=HEAD,size=15,bold=True,color=NAVY); botborder(h,ACCENT,16); return h
def part_page(doc,num,title,icon='',intro=''):
    ip=IMG/f"{icon}.png"
    img_p=doc.add_paragraph(); img_p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    img_p.paragraph_format.page_break_before=True; img_p.paragraph_format.space_before=Pt(115); img_p.paragraph_format.space_after=Pt(10)
    if icon and ip.exists(): img_p.add_run().add_picture(str(ip), width=Mm(34))
    h=doc.add_paragraph(style='Heading 1'); h.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run(h,f"{num}부",font=HEAD,size=13,bold=True,color=ACCENT)
    h2=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=6,after=0); run(h2,title,font=HEAD,size=19,bold=True,color=NAVY)
    if intro:
        pi=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=9,after=0,line=1.6)
        pi.paragraph_format.left_indent=Mm(20); pi.paragraph_format.right_indent=Mm(20)
        run(pi,intro,font=BODY,size=11,italic=True,color=NAVY2)

# ---------- TOC ----------
def add_toc(doc):
    h=heading2(doc,"차례")
    tp=doc.add_paragraph()
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.text="목차를 보려면 [참조 ▸ 목차 업데이트] 또는 F9"
    r.append(t); fld.append(r); tp._p.append(fld)
    s=doc.settings.element
    uf=OxmlElement('w:updateFields'); uf.set(qn('w:val'),'true'); s.insert(0,uf)

# ---------- 빌드 ----------
def build():
    if not (IMG/"heart.png").exists():
        import make_assets; make_assets.build()
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Mm(152); sec.page_height=Mm(225)
    sec.top_margin=Mm(20); sec.bottom_margin=Mm(20); sec.left_margin=Mm(18); sec.right_margin=Mm(18)
    nm=doc.styles['Normal']; nm.font.name=BODY; nm.font.size=Pt(10.5)
    nm.element.rPr.rFonts.set(qn('w:eastAsia'),BODY)
    nm.paragraph_format.line_spacing=1.72; nm.paragraph_format.line_spacing_rule=WD_LINE_SPACING.MULTIPLE
    nm.paragraph_format.space_after=Pt(9)
    _set(nm.element.get_or_add_rPr(),'w:spacing',**{'w:val':'4'})  # 자간 살짝(≈0.2pt)
    for hn,sz in (('Heading 1',18),('Heading 2',15)):
        st=doc.styles[hn]; st.font.name=HEAD; st.font.size=Pt(sz); st.font.bold=True
        st.font.color.rgb=RGBColor.from_string(NAVY)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),HEAD)

    for el in C.parse():
        t=el[0]
        if t=='cover':
            s0=doc.sections[0]
            s0.page_width=Mm(152); s0.page_height=Mm(225)
            s0.top_margin=s0.bottom_margin=s0.left_margin=s0.right_margin=Mm(0)
            cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cp.paragraph_format.space_before=Pt(0); cp.paragraph_format.space_after=Pt(0); cp.paragraph_format.line_spacing=1.0
            cov=IMG/"cover.png"
            if cov.exists(): cp.add_run().add_picture(str(cov), width=Mm(152))
            ns=doc.add_section(WD_SECTION.NEW_PAGE)
            ns.page_width=Mm(152); ns.page_height=Mm(225)
            ns.top_margin=Mm(20); ns.bottom_margin=Mm(20); ns.left_margin=Mm(18); ns.right_margin=Mm(18)
        elif t=='disclaimer':
            heading2(doc,"일러두기")
            c=box(doc,"FFFFFF",NAVY); pp=c.paragraphs[0]; pp.paragraph_format.space_after=Pt(4)
            run(pp,"DISCLAIMER",font=HEAD,size=8.5,bold=True,color=ACCENTD)
            for d in el[1]:
                dp=c.add_paragraph(); dp.paragraph_format.space_after=Pt(4); dp.paragraph_format.line_spacing=1.45
                run(dp,d,size=9.8,color="48433D")
            kp=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=30,after=4); run(kp,"이 책의 핵심은 종목명이 아니라 ‘질문의 구조’입니다.",font=HEAD,size=12.5,bold=True,color=NAVY)
            gp=para(doc,WD_ALIGN_PARAGRAPH.CENTER,after=0,line=1.5); gp.paragraph_format.left_indent=Mm(14); gp.paragraph_format.right_indent=Mm(14)
            run(gp,"이 책은 순서대로 읽어도 좋지만, 6부와 부록은 필요할 때 다시 꺼내보는 실전 노트처럼 활용하셔도 좋습니다.",size=9.6,color=MUTED)
        elif t=='toc': add_toc(doc)
        elif t=='part': part_page(doc, el[1], el[2], el[3] if len(el)>3 else '', el[4] if len(el)>4 else '')
        elif t=='h1big':
            kind=el[1]; icon=el[3] if len(el)>3 else ''
            if kind=='prologue': opener(doc, icon, '', '프롤로그', subtitle=el[2])
            elif kind=='epilogue': opener(doc, icon, '', '에필로그', subtitle=el[2])
            else: opener(doc, icon, 'APPENDIX', el[2])
        elif t=='chapter': opener(doc, el[3] if len(el)>3 else '', f"CHAPTER {el[1]}", f"{el[1]}장. {el[2]}")
        elif t=='keysentence': keysentence(doc, el[1], el[2] if len(el)>2 else '')
        elif t=='ornament': ornament(doc)
        elif t=='figure':
            fp=IMG/f"{el[1]}.png"
            if fp.exists():
                pic=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=3,after=1)
                pic.add_run().add_picture(str(fp), width=Mm(116))
            if len(el)>2 and el[2]:
                cap=para(doc,WD_ALIGN_PARAGRAPH.CENTER,after=6)
                run(cap,"개념도  ",font=HEAD,size=7.5,bold=True,color=ACCENTD); run(cap,el[2],font=HEAD,size=8.5,color=MUTED)
        elif t=='callout': callout(doc, el[1], el[2])
        elif t=='compare': compare(doc, el[1], el[2], el[3], el[4] if len(el)>4 else '')
        elif t=='cards': cards(doc, el[1])
        elif t=='modebar': modebar(doc, el[1])
        elif t=='flow': flow(doc, el[1])
        elif t=='dodont': dodont(doc, el[1], el[2], el[3])
        elif t=='pullquote':
            qm=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=8,after=0,line=1.0); run(qm,"“",font="Georgia",size=30,bold=True,color=ACCENT)
            pq=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=0,after=2,line=1.4); run(pq,el[1],font=BODY,size=14,italic=True,bold=True,color=NAVY)
            rl=para(doc,WD_ALIGN_PARAGRAPH.CENTER,before=1,after=8); rl.paragraph_format.left_indent=Mm(58); rl.paragraph_format.right_indent=Mm(58); botborder(rl,ACCENT,16)
        elif t=='endnote':
            lines=el[1] if isinstance(el[1],(list,tuple)) else [el[1]]
            c=box(doc,"FBF7EF",ACCENT)
            lp=c.paragraphs[0]; lp.paragraph_format.space_after=Pt(3)
            run(lp,"이 장의 정리",font=HEAD,size=8.5,bold=True,color=ACCENTD)
            for i,s in enumerate(lines,1):
                ep=(c.paragraphs[0] if False else c.add_paragraph())
                ep.paragraph_format.space_after=Pt(2); ep.paragraph_format.line_spacing=1.4
                run(ep,f"{i}.  ",font=HEAD,size=10.5,bold=True,color=ACCENTD)
                run(ep,s,font=HEAD,size=10.5,bold=True,color=NAVY)
            spacer(doc)
        elif t=='apxnote':
            c=box(doc,TINT,NAVY2)
            ap=c.paragraphs[0]; ap.paragraph_format.space_after=Pt(0); ap.paragraph_format.line_spacing=1.5
            run(ap,"사용법  ",font=HEAD,size=8,bold=True,color=NAVY2); run(ap,el[1],size=9.6,color="3A4A46"); spacer(doc)
        elif t=='h3':
            p=para(doc,before=6,after=2); run(p,el[1],font=HEAD,size=11.5,bold=True,color=NAVY2)
        elif t=='para':
            body_para(doc, el[1], el[2] if len(el)>2 else False)
        elif t=='bullets':
            for it in el[1]:
                bp=para(doc,after=2,indent=4); run(bp,"· ",bold=True,color=ACCENT); run(bp,it,size=10,color=INK)
        elif t=='numlist':
            for num,it in el[1]:
                bp=para(doc,after=2,indent=4); run(bp,f"{num}. ",bold=True,color=ACCENTD); run(bp,it,size=10,color=INK)
        elif t=='writeline':
            wl=para(doc,before=1,after=3); wl.paragraph_format.left_indent=Mm(2); wl.paragraph_format.right_indent=Mm(2)
            botborder(wl, "C9C2B4", 6)
        elif t=='cat':
            cp=para(doc,before=8,after=3); botborder(cp,ACCENT,12)
            run(cp,f" {el[1]}  ",font=HEAD,size=9,bold=True,color="FFFFFF")
            run(cp,el[2],font=HEAD,size=11.5,bold=True,color=NAVY)
        elif t=='q':
            qmode,num,text=el[1],el[2],el[3]
            qp=para(doc,before=5,after=2)
            run(qp,("Q"+num+" " if qmode else num+". "),font=HEAD,size=10.5,bold=True,color=ACCENTD)
            run(qp,text,font=HEAD,size=10.5,bold=True,color=INK)
        elif t=='prompt':
            c=box(doc,PROMPTBG,NAVY2); tp=c.paragraphs[0]; tp.paragraph_format.space_after=Pt(2)
            run(tp,"붙여넣을 프롬프트",font=HEAD,size=7.8,bold=True,color=NAVY2)
            bp=c.add_paragraph(); bp.paragraph_format.space_after=Pt(0); bp.paragraph_format.line_spacing=1.4
            run(bp,el[1],size=9.6,color="33433F"); spacer(doc,1)
        elif t=='fill':
            fp=para(doc,before=2,after=2); run(fp,el[1]+" ",font=HEAD,size=9,color=MUTED); botborder(fp,LINE,6)

    doc.save(str(OUT)); print("[ok] DOCX →", OUT)

if __name__=="__main__": build()
